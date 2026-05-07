from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from kaoqin.generator import (
    DEFAULT_DINGTALK_SHEET,
    generate_attendance,
    get_day_columns_from_dingtalk,
    is_dingtalk_sheet,
    resolve_dingtalk_sheet,
    resolve_input_file,
)
def create_dingtalk_workbook(path: Path, sheet_name: str = DEFAULT_DINGTALK_SHEET):
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet_name
    ws["A1"] = "月度汇总 统计日期：2026-03-01 至 2026-03-31"
    ws["A2"] = "报表生成时间：2026-04-04 11:04"
    ws["A3"] = "姓名"
    ws["B3"] = "考勤结果"
    ws["B4"] = "日"
    ws["C4"] = "2"
    ws["D4"] = "3"
    ws["A5"] = "刘晓"
    ws["B5"] = "休息\n(-)"
    ws["C5"] = "正常班:正常\n(08:17,18:30)"
    ws["D5"] = "正常班:病假\n(-)"
    ws["A6"] = "袁晨"
    ws["B6"] = "休息\n(-)"
    ws["C6"] = "正常班:正常\n(08:17,18:30)"
    ws["D6"] = "正常班:事假\n(-)"
    workbook.save(path)


def build_row_index(ws):
    index = {}
    current_name = None
    for row in range(1, ws.max_row + 1):
        name = ws.cell(row, 1).value
        att_type = ws.cell(row, 2).value
        if name and name not in {"姓  名", "邮政客户部", "电力技术部"}:
            current_name = str(name).strip()
            index[current_name] = {}
        if current_name and att_type:
            index[current_name][str(att_type).strip()] = row
    return index


def test_is_dingtalk_sheet_detects_expected_layout(tmp_path):
    workbook_path = tmp_path / "sample.xlsx"
    create_dingtalk_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    assert is_dingtalk_sheet(workbook[DEFAULT_DINGTALK_SHEET]) is True


def test_resolve_dingtalk_sheet_finds_non_default_sheet(tmp_path):
    workbook_path = tmp_path / "sample.xlsx"
    create_dingtalk_workbook(workbook_path, sheet_name="SheetX")
    workbook = load_workbook(workbook_path)
    assert resolve_dingtalk_sheet(workbook) == "SheetX"


def test_resolve_input_file_auto_detects_data_file(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    workbook_path = data_dir / "3月考勤电力，邮政.xlsx"
    create_dingtalk_workbook(workbook_path)

    resolved = resolve_input_file(None, None, base_dir=tmp_path)
    assert resolved == str(workbook_path)


def test_generate_attendance_creates_workbook_with_formulas_and_alias(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    workbook_path = data_dir / "3月考勤电力，邮政.xlsx"
    create_dingtalk_workbook(workbook_path)

    generate_attendance(config_dir=str(tmp_path), reference_date=date(2026, 4, 15))

    output_path = data_dir / "邮政+电力（3月）_自动生成.xlsx"
    assert output_path.exists()

    workbook = load_workbook(output_path, data_only=False)
    ws = workbook["考勤（一）"]
    assert ws["A1"].value == "2026年3月份考勤汇总统计表"
    row_index = build_row_index(ws)

    liuxiao_work_row = row_index["刘晓"]["出勤 √"]
    liuxiao_sick_row = row_index["刘晓"]["病假 ※"]
    yuanchen_work_row = row_index["袁辰"]["出勤 √"]
    yuanchen_personal_row = row_index["袁辰"]["事假 △"]

    assert ws.cell(liuxiao_work_row, 34).value.startswith("=(")
    assert ws.cell(liuxiao_sick_row, 5).value == 8
    assert ws.cell(yuanchen_work_row, 34).value.startswith("=(")
    assert ws.cell(yuanchen_personal_row, 5).value == 8


def test_get_day_columns_from_dingtalk_handles_weekend_headers_without_overwriting_day_one():
    workbook = Workbook()
    ws = workbook.active
    ws["B4"] = "1"
    ws["C4"] = "2"
    ws["D4"] = "3"
    ws["E4"] = "六"
    ws["F4"] = "日"
    ws["G4"] = "6"

    day_cols = get_day_columns_from_dingtalk(ws)

    assert day_cols[1] == 2
    assert day_cols[2] == 3
    assert day_cols[3] == 4
    assert day_cols[6] == 7


def test_get_day_columns_from_dingtalk_supports_legacy_first_day_header():
    workbook = Workbook()
    ws = workbook.active
    ws["B4"] = "日"
    ws["C4"] = "2"
    ws["D4"] = "3"

    day_cols = get_day_columns_from_dingtalk(ws)

    assert day_cols[1] == 2
    assert day_cols[2] == 3
    assert day_cols[3] == 4
