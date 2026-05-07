from __future__ import annotations

import argparse
import calendar
import logging
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import load_json_config
from .exceptions import ConfigError, GenerationError, InputFileError, KaoqinError, SheetDetectionError
from .logging_utils import setup_logging


DEFAULT_DINGTALK_SHEET = "月度汇总"
DEFAULT_DATA_DIRNAME = "data"
logger = logging.getLogger(__name__)

FONT = Font(name="宋体", size=9)
TITLE_FONT = Font(name="宋体", size=9, bold=True)
HEADER_FONT = Font(name="黑体", size=9, bold=True)
FIRST_COLUMN_FONT = Font(name="黑体", size=9, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="000000")
HAIR = Side(style="hair", color="000000")
DAY_FILL = PatternFill(fill_type="solid", fgColor="FF00B0F0")
TYPE_FILL = PatternFill(fill_type="solid", fgColor="FFF2F2F2")


@dataclass(frozen=True)
class RuntimeConfig:
    holiday_calendars: dict[str, dict[str, set[str]]]
    department_layout: list[tuple[str, list[tuple[str, list[str]]]]]
    auto_workday_people: set[str]
    name_alias: dict[str, str]
    status_to_type: dict[str, tuple]
    attendance_type_meta: dict[str, tuple[str, str]]


def load_runtime_config(base_dir: Path | None = None) -> RuntimeConfig:
    try:
        holiday_raw = load_json_config("holiday_calendars", base_dir=base_dir)
        attendance_layout_raw = load_json_config("attendance_layout", base_dir=base_dir)
        attendance_config_raw = load_json_config("attendance_config", base_dir=base_dir)
        attendance_rules_raw = load_json_config("attendance_rules", base_dir=base_dir)
    except ConfigError:
        raise
    except Exception as exc:
        raise ConfigError("加载项目配置失败") from exc

    holiday_calendars = {
        str(year): {
            "holidays": set(config.get("holidays", [])),
            "workdays": set(config.get("workdays", [])),
        }
        for year, config in holiday_raw.items()
    }
    department_layout = []
    for department in attendance_layout_raw.get("departments", []):
        employees = []
        for employee in department.get("employees", []):
            employees.append((employee["name"], employee.get("attendance_types", [])))
        department_layout.append((department["name"], employees))

    status_to_type = {
        status: tuple(values)
        for status, values in attendance_rules_raw.get("status_to_type", {}).items()
    }
    attendance_type_meta = {
        att_type: tuple(values)
        for att_type, values in attendance_rules_raw.get("attendance_type_meta", {}).items()
    }

    return RuntimeConfig(
        holiday_calendars=holiday_calendars,
        department_layout=department_layout,
        auto_workday_people=set(attendance_config_raw.get("auto_workday_people", [])),
        name_alias=attendance_config_raw.get("name_alias", {}),
        status_to_type=status_to_type,
        attendance_type_meta=attendance_type_meta,
    )


def is_workday(runtime_config: RuntimeConfig, year: int, month: int, day: int) -> bool:
    current_date = date(year, month, day)
    current_str = current_date.isoformat()
    holiday_calendar = runtime_config.holiday_calendars.get(str(year))
    if holiday_calendar:
        if current_str in holiday_calendar["workdays"]:
            return True
        if current_str in holiday_calendar["holidays"]:
            return False
    return current_date.weekday() < 5


def parse_period_from_dingtalk(ws) -> tuple[int, int]:
    title = str(ws.cell(1, 1).value or "")
    match = re.search(r"统计日期[:：]\s*(\d{4})-(\d{2})-\d{2}\s*至\s*(\d{4})-(\d{2})-\d{2}", title)
    if not match:
        raise ValueError("无法从钉钉表头识别统计月份")

    start_year = int(match.group(1))
    start_month = int(match.group(2))
    end_year = int(match.group(3))
    end_month = int(match.group(4))
    if (start_year, start_month) != (end_year, end_month):
        raise ValueError("钉钉文件跨月份，当前脚本只支持单月统计")
    return start_year, start_month


def get_target_period(reference_date: date | None = None) -> tuple[int, int]:
    current_date = reference_date or date.today()
    if current_date.month == 1:
        return current_date.year - 1, 12
    return current_date.year, current_date.month - 1


def get_output_file(year: int, month: int) -> str:
    return f"邮政+电力（{month}月）_自动生成.xlsx"


def parse_args():
    parser = argparse.ArgumentParser(description="根据钉钉月度汇总生成考勤汇总表")
    parser.add_argument("input_file", nargs="?", help="钉钉导出的月度汇总 Excel 文件；不传时自动扫描 data 目录")
    parser.add_argument("-s", "--sheet", help="钉钉工作表名称；不传时自动识别，优先使用月度汇总")
    parser.add_argument("-o", "--output", help="输出文件名；不传时按月份自动生成到 data 目录")
    parser.add_argument("--config-dir", help="配置文件目录；不传时优先读取当前目录下的 JSON")
    parser.add_argument("--log-level", default="INFO", help="日志级别，默认 INFO")
    parser.add_argument("--log-file", help="日志文件路径；不传时默认写入 logs/kaoqin.log")
    return parser.parse_args()


def get_data_dir(base_dir: Path | None = None) -> Path:
    root_dir = base_dir or Path.cwd()
    return root_dir / DEFAULT_DATA_DIRNAME


def resolve_path(path_str: str, base_dir: Path | None = None) -> Path:
    path = Path(path_str)
    if path.is_absolute():
        return path
    root_dir = base_dir or Path.cwd()
    return root_dir / path


def list_candidate_input_files(base_dir: Path | None = None) -> list[Path]:
    current_dir = get_data_dir(base_dir=base_dir)
    if not current_dir.exists():
        return []
    files = []
    for path in current_dir.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        if "自动生成" in path.stem:
            continue
        files.append(path)

    def sort_key(path: Path):
        # Prefer likely DingTalk source files over unrelated workbooks.
        preferred = 0 if "考勤" in path.stem else 1
        return (preferred, -path.stat().st_mtime, path.name)

    return sorted(files, key=sort_key)


def is_dingtalk_sheet(ws) -> bool:
    # DingTalk exports are identified from a small set of stable header cells.
    title = str(ws.cell(1, 1).value or "")
    header_a3 = str(ws.cell(3, 1).value or "").strip()
    header_b3 = str(ws.cell(3, 2).value or "").strip()
    header_b4 = str(ws.cell(4, 2).value or "").strip()
    has_period = "统计日期" in title
    has_headers = header_a3 == "姓名" and header_b3 == "考勤结果"
    has_day_header = header_b4 in {"日", "1"}
    return has_period and has_headers and has_day_header


def resolve_dingtalk_sheet(workbook, sheet_name: str | None = None) -> str:
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise SheetDetectionError(f"找不到工作表：{sheet_name}")
        return sheet_name

    if DEFAULT_DINGTALK_SHEET in workbook.sheetnames:
        ws = workbook[DEFAULT_DINGTALK_SHEET]
        if is_dingtalk_sheet(ws):
            return DEFAULT_DINGTALK_SHEET

    for candidate_name in workbook.sheetnames:
        ws = workbook[candidate_name]
        if is_dingtalk_sheet(ws):
            return candidate_name

    raise SheetDetectionError("无法自动识别钉钉月度汇总工作表，请通过 -s/--sheet 指定")


def resolve_input_file(input_file: str | None, dingtalk_sheet: str | None = None, base_dir: Path | None = None) -> str:
    if input_file:
        resolved = resolve_path(input_file, base_dir=base_dir)
        if not resolved.exists():
            raise InputFileError(f"输入文件不存在：{resolved}")
        return str(resolved)

    for candidate in list_candidate_input_files(base_dir=base_dir):
        try:
            workbook = load_workbook(candidate, read_only=True, data_only=True)
            resolve_dingtalk_sheet(workbook, dingtalk_sheet)
            logger.info("Auto detected input file: %s", candidate)
            return str(candidate)
        except Exception as exc:
            logger.debug("Skip candidate file %s: %s", candidate, exc)
            continue

    raise InputFileError("data 目录下找不到可用的钉钉月度汇总 Excel 文件，请显式传入 input_file")


def get_day_columns_from_dingtalk(ws) -> dict[int, int]:
    day_cols = {}
    for col in range(2, ws.max_column + 1):
        value = ws.cell(4, col).value
        if value is None:
            continue
        value = str(value).strip()
        if value == "日":
            day = 1
        elif value.isdigit():
            day = int(value)
        else:
            continue
        if 1 <= day <= 31:
            day_cols[day] = col
    return day_cols


def parse_status(text) -> str:
    if text is None:
        return "空"
    text = str(text)
    checks = [
        ("休息", "休息"),
        ("旷工", "旷工"),
        ("迟到", "迟到"),
        ("项目休假", "项目休假"),
        ("陪产假", "陪产假"),
        ("产假", "产假"),
        ("育儿假", "育儿假"),
        ("婚假", "婚假"),
        ("丧假", "丧假"),
        ("病假", "病假"),
        ("事假", "事假"),
        ("年假", "年假"),
        ("倒休", "倒休"),
        ("出差", "出差"),
    ]
    for keyword, status in checks:
        if keyword in text:
            return status
    if any(keyword in text for keyword in ["外勤", "缺卡", "补卡", "早退", "正常"]):
        return "正常"
    return "未知"


def make_border(left=THIN, right=THIN, top=THIN, bottom=THIN):
    return Border(left=left, right=right, top=top, bottom=bottom)


def get_day_fill(runtime_config: RuntimeConfig, year: int, month: int, day: int):
    if is_workday(runtime_config, year, month, day):
        return None
    return DAY_FILL


def build_sheet(runtime_config: RuntimeConfig, year: int, month: int):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "考勤（一）"
    days_in_month = calendar.monthrange(year, month)[1]
    total_col = days_in_month + 3
    unit_col = total_col + 1
    category_col = total_col + 2
    last_col = category_col
    last_letter = get_column_letter(last_col)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A1"] = f"{year}年{month}月份考勤汇总统计表"
    ws["A2"] = "注:带*号黄色阴影为新入职人员;灰色阴影部分为离职人员"
    ws["A1"].font = HEADER_FONT
    ws["A2"].font = HEADER_FONT
    ws["A1"].alignment = CENTER
    ws["A2"].alignment = LEFT
    ws["A1"].border = make_border()
    ws["A2"].border = make_border()

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 9.5
    for col in range(3, total_col):
        ws.column_dimensions[get_column_letter(col)].width = 4.2
    ws.column_dimensions[get_column_letter(total_col)].width = 8
    ws.column_dimensions[get_column_letter(unit_col)].width = 5.2
    ws.column_dimensions[get_column_letter(category_col)].width = 8.5

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 22.5
    ws.row_dimensions[3].height = 23

    ws.cell(3, 1).value = "姓  名"
    ws.cell(3, 2).value = 0
    for day in range(1, days_in_month + 1):
        ws.cell(3, day + 2).value = f"{day}日"
    ws.cell(3, total_col).value = "合计"
    ws.cell(3, unit_col).value = "单位"
    ws.cell(3, category_col).value = "类型"

    for col in range(1, last_col + 1):
        cell = ws.cell(3, col)
        cell.font = TITLE_FONT if col in {total_col, unit_col, category_col} else FONT
        if col == 1:
            cell.font = FIRST_COLUMN_FONT
        cell.alignment = CENTER
        if 3 <= col < total_col:
            fill = get_day_fill(runtime_config, year, month, col - 2)
            if fill:
                cell.fill = fill
        cell.border = make_border()
    ws.cell(3, total_col).number_format = '0.00"天"'

    row = 4
    row_index = {}
    for department, employees in runtime_config.department_layout:
        # Each department gets a header row, followed by a variable-height block per person.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        dept_cell = ws.cell(row, 1)
        dept_cell.value = department
        dept_cell.font = FIRST_COLUMN_FONT
        dept_cell.alignment = LEFT
        dept_cell.border = make_border()
        for col in range(3, last_col + 1):
            cell = ws.cell(row, col)
            cell.border = make_border()
            if col < total_col:
                fill = get_day_fill(runtime_config, year, month, col - 2)
                if fill:
                    cell.fill = fill
        row += 1

        for name, att_types in employees:
            start_row = row
            end_row = row + len(att_types) - 1
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            name_cell = ws.cell(start_row, 1)
            name_cell.value = name
            name_cell.font = FIRST_COLUMN_FONT
            name_cell.alignment = CENTER

            row_index[name] = {}
            for offset, att_type in enumerate(att_types):
                current_row = start_row + offset
                row_index[name][att_type] = current_row
                a_cell = ws.cell(current_row, 1)
                b_cell = ws.cell(current_row, 2)
                total_cell = ws.cell(current_row, total_col)
                unit_cell = ws.cell(current_row, unit_col)
                category_cell = ws.cell(current_row, category_col)

                a_cell.border = make_border(left=THIN, right=HAIR, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)
                b_cell.value = att_type
                b_cell.font = FONT
                b_cell.alignment = CENTER
                b_cell.fill = TYPE_FILL
                b_cell.border = make_border(left=HAIR, right=HAIR, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)

                for col in range(3, total_col):
                    cell = ws.cell(current_row, col)
                    cell.font = FONT
                    cell.alignment = CENTER
                    fill = get_day_fill(runtime_config, year, month, col - 2)
                    if fill:
                        cell.fill = fill
                    cell.border = make_border(left=HAIR, right=HAIR, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)

                total_cell.font = FONT
                total_cell.alignment = CENTER
                total_cell.border = make_border(left=HAIR, right=THIN, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)
                if att_type == "出勤 √":
                    total_cell.number_format = '0.00"天"'

                unit_value, category_value = runtime_config.attendance_type_meta.get(att_type, ("", ""))
                unit_cell.value = unit_value
                unit_cell.font = FONT
                unit_cell.alignment = CENTER
                unit_cell.border = make_border(left=THIN, right=THIN, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)

                category_cell.value = category_value
                category_cell.font = FONT
                category_cell.alignment = CENTER
                category_cell.border = make_border(left=THIN, right=THIN, top=THIN if offset == 0 else HAIR, bottom=THIN if current_row == end_row else HAIR)

            row = end_row + 1

    ws.freeze_panes = "C4"
    ws.print_area = f"A1:{last_letter}{row - 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    return workbook, ws, row_index, total_col


def fill_auto_workdays(runtime_config: RuntimeConfig, ws, row_index, year: int, month: int):
    days = calendar.monthrange(year, month)[1]
    for name in runtime_config.auto_workday_people:
        work_row = row_index.get(name, {}).get("出勤 √")
        if not work_row:
            logger.warning("布局中找不到 %s 的 出勤 行", name)
            continue
        for day in range(1, days + 1):
            if is_workday(runtime_config, year, month, day):
                ws.cell(work_row, day + 2).value = 8


def write_status(runtime_config: RuntimeConfig, ws, row_index, name: str, day: int, status: str):
    if name not in row_index:
        logger.warning("布局中找不到人员：%s", name)
        return
    if name in runtime_config.auto_workday_people and status == "正常":
        return
    if status == "正常":
        work_row = row_index[name].get("出勤 √")
        if work_row:
            ws.cell(work_row, day + 2).value = 8
        return
    if status in {"休息", "空"}:
        return

    mapping = runtime_config.status_to_type.get(status)
    if not mapping:
        logger.warning("未知状态：%s %s日 %s", name, day, status)
        return
    for idx in range(0, len(mapping), 2):
        att_type = mapping[idx]
        value = mapping[idx + 1]
        row = row_index[name].get(att_type)
        if row:
            ws.cell(row, day + 2).value = value


def update_total(runtime_config: RuntimeConfig, ws, row_index, days_in_month: int, total_col: int):
    start_letter = get_column_letter(3)
    end_letter = get_column_letter(days_in_month + 2)
    for rows in row_index.values():
        for att_type, row in rows.items():
            total_cell = ws.cell(row, total_col)
            day_range = f"{start_letter}{row}:{end_letter}{row}"
            if att_type == "出勤 √":
                # Attendance is reported in days, so all hour-based rows are folded in and divided by 8.
                total_parts = [f"SUM({day_range})"]
                for other_type, other_row in rows.items():
                    if other_type == "出勤 √":
                        continue
                    unit_value, _ = runtime_config.attendance_type_meta.get(other_type, ("", ""))
                    if unit_value == "小时":
                        total_parts.append(f"SUM({start_letter}{other_row}:{end_letter}{other_row})")
                total_cell.value = f"=({'+'.join(total_parts)})/8"
            else:
                total_cell.value = f"=SUM({day_range})"


def generate_attendance(
    input_file: str | None = None,
    dingtalk_sheet: str | None = None,
    output_file: str | None = None,
    config_dir: str | None = None,
    reference_date: date | None = None,
):
    base_dir = Path(config_dir).resolve() if config_dir else Path.cwd()
    try:
        runtime_config = load_runtime_config(base_dir=base_dir)
        input_file = resolve_input_file(input_file, dingtalk_sheet, base_dir=base_dir)
        logger.info("Use input file: %s", input_file)

        ding_wb = load_workbook(input_file, data_only=True)
        resolved_sheet = resolve_dingtalk_sheet(ding_wb, dingtalk_sheet)
        ding_ws = ding_wb[resolved_sheet]
        logger.info("Use worksheet: %s", resolved_sheet)

        source_year, source_month = parse_period_from_dingtalk(ding_ws)
        year, month = get_target_period(reference_date=reference_date)
        days_in_month = calendar.monthrange(year, month)[1]
        logger.info("Target period: %s-%02d", year, month)
        if (source_year, source_month) != (year, month):
            logger.warning(
                "DingTalk source period %s-%02d does not match target period %s-%02d",
                source_year,
                source_month,
                year,
                month,
            )

        if output_file is None:
            data_dir = get_data_dir(base_dir=base_dir)
            data_dir.mkdir(parents=True, exist_ok=True)
            output_file = str(data_dir / get_output_file(year, month))
        else:
            output_file = str(resolve_path(output_file, base_dir=base_dir))
        logger.info("Output file: %s", output_file)

        workbook, ws, row_index, total_col = build_sheet(runtime_config, year, month)
        fill_auto_workdays(runtime_config, ws, row_index, year, month)
        day_cols = get_day_columns_from_dingtalk(ding_ws)

        for row in range(5, ding_ws.max_row + 1):
            raw_name = ding_ws.cell(row, 1).value
            if raw_name is None:
                continue
            name = runtime_config.name_alias.get(str(raw_name).strip(), str(raw_name).strip())
            for day, ding_col in day_cols.items():
                if day > days_in_month:
                    continue
                status = parse_status(ding_ws.cell(row, ding_col).value)
                write_status(runtime_config, ws, row_index, name, day, status)

        update_total(runtime_config, ws, row_index, days_in_month, total_col)
        workbook.save(output_file)
        logger.info("Attendance workbook generated successfully: %s", output_file)
        print(f"已生成：{output_file}")
    except KaoqinError:
        raise
    except Exception as exc:
        raise GenerationError("生成考勤汇总文件失败") from exc


def main():
    args = parse_args()
    base_dir = Path(args.config_dir).resolve() if args.config_dir else Path.cwd()
    log_path = setup_logging(base_dir=base_dir, level=args.log_level, log_file=args.log_file)
    logger.info("Logging to: %s", log_path)
    try:
        generate_attendance(
            input_file=args.input_file,
            dingtalk_sheet=args.sheet,
            output_file=args.output,
            config_dir=args.config_dir,
        )
    except KaoqinError as exc:
        logger.exception("Kaoqin failed: %s", exc)
        print(f"错误：{exc}")
        raise SystemExit(1) from exc
    except Exception as exc:
        logger.exception("Unexpected error")
        print(f"未预期错误：{exc}")
        raise SystemExit(1) from exc
